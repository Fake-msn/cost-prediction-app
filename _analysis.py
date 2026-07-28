# -*- coding: utf-8 -*-
import os, sys, json
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl
from collections import Counter

base = r'd:\工程造价预测AI项目模板'
dirs = sorted([d for d in os.listdir(base) if os.path.isdir(os.path.join(base, d))])
print(f'Total project dirs: {len(dirs)}')

# Find project 7
target_dir = None
for d in dirs:
    if '莲花' in d:
        target_dir = d
        break
if not target_dir and len(dirs) > 6:
    target_dir = dirs[6]  # 7th directory (0-indexed)

print(f'Target dir: {target_dir}')

target_path = os.path.join(base, target_dir)
xlsx_files = [f for f in os.listdir(target_path) if f.endswith('.xlsx') and not f.startswith('~')]
print(f'xlsx files: {xlsx_files}')

main_xlsx = [f for f in xlsx_files if '项目信息' not in f]
if main_xlsx:
    wb_path = os.path.join(target_path, main_xlsx[0])
    wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=True)
    print(f'\nTotal sheets: {len(wb.sheetnames)}')
    
    # Categorize
    patterns = ['封-2','扉-2','表-01','表-02','表-03','表-04','表-08','表-11','表-12-1','表-12-3','表-12-5','表-12','表-13','表-21']
    types = Counter()
    for name in wb.sheetnames:
        matched = False
        for p in patterns:
            if p in name:
                types[p] += 1
                matched = True
                break
        if not matched:
            types['other'] += 1
    
    print('\nSheet type counts:')
    for t, c in sorted(types.items()):
        print(f'  {t}: {c} sheets')
    
    # Print all sheet names
    print('\nAll sheet names:')
    for i, name in enumerate(wb.sheetnames):
        print(f'  [{i}] {name[:80]}')
    
    wb.close()