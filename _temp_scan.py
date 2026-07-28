import openpyxl,os,json
source=r'd:\工程造价预测AI项目模板'
results=[]
for folder in sorted(os.listdir(source)):
    p=os.path.join(source,folder)
    if not os.path.isdir(p): continue
    xlsx=[f for f in os.listdir(p) if f.endswith('.xlsx') and not f.startswith('~')]
    has_info=any('项目信息' in f for f in xlsx)
    main=[f for f in xlsx if '项目信息' not in f]
    sc=0;t02=t04=t08=t21=False
    if main:
        try:
            wb=openpyxl.load_workbook(os.path.join(p,main[0]),read_only=True)
            sc=len(wb.sheetnames);sn=' '.join(wb.sheetnames)
            t02='表-02' in sn;t04='表-04' in sn;t08='表-08' in sn;t21='表-21' in sn
            wb.close()
        except: pass
    results.append({'folder':folder,'sheets':sc,'has_info':has_info,'t02':t02,'t04':t04,'t08':t08,'t21':t21,'xlsx_count':len(xlsx)})
with open(r'c:\Users\aaa12\Documents\Qoder\2026-07-27\chat-2\cost-prediction-app\_temp_scan_result.json','w',encoding='utf-8') as out:
    json.dump(results,out,ensure_ascii=False,indent=2)
print(f'Scanned {len(results)} projects')
