# -*- coding: utf-8 -*-
import os, sys
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

base = r'd:\工程造价预测AI项目模板'
target_dir = '7莲花C地块拆迁统建还房1~2#楼工程'
target_path = os.path.join(base, target_dir)
xlsx_files = [f for f in os.listdir(target_path) if f.endswith('.xlsx') and not f.startswith('~$') and '项目信息' not in f]
wb_path = os.path.join(target_path, xlsx_files[0])
wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=True)

# Analyze 表-08 for categorization into 基础/主体/屋面/外墙
print('=== 表-08: 分部分项工程 - Section headers and sample items ===')
t08_sheets = [s for s in wb.sheetnames if '表-08' in s]
ws = wb[t08_sheets[0]]
sections = []
for row in ws.iter_rows(values_only=True):
    if row and len(row) >= 3:
        name = str(row[2]).strip() if row[2] else ''
        if name and ('工程' in name or '分部' in name or name in ['小计','合计']) and len(name) < 30:
            sections.append(name)
            print(f'  Section: {name}')

print(f'\nTotal sections in first 表-08: {len(sections)}')

# Analyze 表-13 for 规费 breakdown
print('\n=== 表-13: 规费、税金 breakdown ===')
t13_sheets = [s for s in wb.sheetnames if '表-13' in s]
ws = wb[t13_sheets[0]]
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if row and len(row) >= 3:
        name = str(row[1]).strip() if row[1] else ''
        val = row[5] if len(row) > 5 else None
        if name:
            print(f'  {name}: {val}')

# Analyze 表-04 for cost breakdown
print('\n=== 表-04: 单位工程汇总 (first sheet) ===')
t04_sheets = [s for s in wb.sheetnames if '表-04' in s]
ws = wb[t04_sheets[0]]
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if row and len(row) >= 3:
        content = str(row[1]).strip() if row[1] else ''
        val = row[2] if len(row) > 2 else None
        if content:
            print(f'  {content}: {val}')

# Check 表-08 for 定额人工费 column
print('\n=== 表-08: Checking 定额人工费 column ===')
ws = wb[t08_sheets[0]]
for i, row in enumerate(ws.iter_rows(max_row=10, values_only=True)):
    if row:
        print(f'  Row {i}: {[str(c)[:20] if c else "" for c in row]}')

# Sum 定额人工费 from first 表-08
total_labor = 0
ws = wb[t08_sheets[0]]
for row in ws.iter_rows(values_only=True):
    if row and len(row) >= 9:
        val = row[8]  # 定额人工费 column
        if val and isinstance(val, (int, float)):
            total_labor += val
print(f'\nTotal 定额人工费 from first 表-08 sheet: {total_labor:,.2f}')

# Sum across all 表-08 sheets
total_labor_all = 0
for sheet_name in t08_sheets:
    ws = wb[sheet_name]
    for row in ws.iter_rows(values_only=True):
        if row and len(row) >= 9:
            val = row[8]
            if val and isinstance(val, (int, float)):
                total_labor_all += val
print(f'Total 定额人工费 from ALL 表-08 sheets: {total_labor_all:,.2f}')

wb.close()