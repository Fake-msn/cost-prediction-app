# -*- coding: utf-8 -*-
import os, sys
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

base = r'd:\工程造价预测AI项目模板'
target_dir = '7莲花C地块拆迁统建还房1~2#楼工程'
target_path = os.path.join(base, target_dir)
xlsx_files = [f for f in os.listdir(target_path) if f.endswith('.xlsx') and not f.startswith('~$') and '项目信息' not in f]
print(f'Found xlsx files: {xlsx_files}')
wb_path = os.path.join(target_path, xlsx_files[0])
print(f'Loading: {wb_path}')
wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=True)

patterns = ['封-2','扉-2','表-01','表-02','表-03','表-04','表-08','表-11','表-12-1','表-12-3','表-12-5','表-12','表-13','表-21']
seen = set()
for name in wb.sheetnames:
    sheet_type = None
    for p in patterns:
        if p in name:
            sheet_type = p
            break
    if sheet_type and sheet_type not in seen:
        seen.add(sheet_type)
        ws = wb[name]
        print(f'\n=== {sheet_type} ({name[:70]}) ===')
        for i, row in enumerate(ws.iter_rows(max_row=15, values_only=True)):
            row_data = list(row)
            print(f'  Row {i}: {row_data}')

wb.close()