"""excel_reader 真实 DuckDB 引擎集成测试（非 mock）。

覆盖 POC 验证过的行为：
- read_xlsx 快路径返回真实 DataFrame；
- stop_at_empty=false 下含中部空行的文件不被截断（读全）；
- 合并单元格仅左上角有值、其余为 NULL（记录该限制，BOQ 须走 openpyxl）；
- COPY ... TO xlsx 导出往返一致；
- 2 万行完整读取（无截断）；
- data_loader.import_from_excel 端到端：经快路径读取 → 现有清洗 → 入库，异常行被过滤。

若 excel 扩展无法加载（无网络 INSTALL）则整体 skip，避免 CI 抖动。
所有用例在 tempfile 独立库上运行，绝不触碰 data/cost_prediction.duckdb。
"""
import os
import sys
import tempfile

import pandas as pd
import pytest
from openpyxl import Workbook

_HERE = os.path.dirname(os.path.abspath(__file__))
_BACKEND = os.path.dirname(_HERE)
_SCRIPTS = os.path.join(_BACKEND, "scripts")
for _p in (_BACKEND, _SCRIPTS):
    if _p not in sys.path:
        sys.path.insert(0, _p)

import duckdb  # noqa: E402
import init_duckdb  # noqa: E402
from data_loader import DataLoader  # noqa: E402
from excel_reader import (  # noqa: E402
    export_dataframe_to_xlsx,
    read_table,
    read_xlsx_to_dataframe,
)

REQUIRED = ["项目名称", "建筑类型", "结构类型", "总建筑面积", "楼层数",
            "所在地区", "建造年份", "装修标准", "项目总造价", "单方造价"]


def _excel_available() -> bool:
    try:
        con = duckdb.connect()
        try:
            con.execute("LOAD excel;")
        except Exception:
            con.execute("INSTALL excel;")
            con.execute("LOAD excel;")
        con.close()
        return True
    except Exception:
        return False


pytestmark = pytest.mark.skipif(
    not _excel_available(),
    reason="DuckDB excel 扩展不可用（无网络 INSTALL excel）",
)


def _save_wb(wb) -> bytes:
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf:
        path = tf.name
    wb.save(path)
    with open(path, "rb") as f:
        content = f.read()
    os.unlink(path)
    return content


def test_read_xlsx_fast_path_returns_dataframe():
    wb = Workbook()
    ws = wb.active
    ws.append(["项目名称", "单方造价"])
    ws.append(["A", 4500.0])
    ws.append(["B", 6200.0])
    content = _save_wb(wb)

    df = read_xlsx_to_dataframe(content, "fast.xlsx")
    assert isinstance(df, pd.DataFrame)
    assert list(df.columns) == ["项目名称", "单方造价"]
    assert len(df) == 2


def test_stop_at_empty_false_reads_through_blank_rows():
    """中部空行不得截断：最后一行的哨兵值必须被读到。"""
    wb = Workbook()
    ws = wb.active
    ws.append(["项目名称", "单方造价"])
    ws.append(["row1", 1000.0])
    ws.append([None, None])          # 中部空行
    ws.append([None, None])
    ws.append(["SENTINEL_LAST", 2000.0])
    content = _save_wb(wb)

    df = read_xlsx_to_dataframe(content, "blanks.xlsx")
    names = set(df["项目名称"].dropna().astype(str))
    assert "SENTINEL_LAST" in names, "含空行文件被截断，未读到末尾数据"


def test_merged_cells_only_top_left_has_value():
    """合并单元格：read_xlsx 仅左上角有值，其余为 NULL（BOQ 须保留 openpyxl 展开）。"""
    wb = Workbook()
    ws = wb.active
    ws.append(["分部名称", "项目编码"])
    ws.append(["土石方工程", "010101001001"])
    ws.append([None, "010101002001"])
    ws.append([None, "010103001001"])
    ws.merge_cells("A2:A4")
    content = _save_wb(wb)

    df = read_xlsx_to_dataframe(content, "merged.xlsx")
    div = df["分部名称"].tolist()
    assert str(div[0]) == "土石方工程"
    # 其余合并覆盖行为空（NaN/None）——结构信息丢失，无法程序化展开
    assert all(pd.isna(v) or v in (None, "") for v in div[1:])


def test_export_copy_to_xlsx_roundtrip():
    df = pd.DataFrame({
        "project_id": ["P1", "P2"],
        "name": ["住宅", "办公"],
        "total_cost": [1.5e8, 3.6e8],
    })
    with tempfile.TemporaryDirectory() as d:
        dest = os.path.join(d, "report.xlsx")
        out = export_dataframe_to_xlsx(df, dest)
        assert os.path.exists(out) and os.path.getsize(out) > 0
        back = read_xlsx_to_dataframe(out, "report.xlsx")
        assert len(back) == 2
        assert set(back["project_id"].astype(str)) == {"P1", "P2"}
        assert abs(float(back["total_cost"].sum()) - (1.5e8 + 3.6e8)) < 1.0


def test_read_table_20k_rows_no_truncation():
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("S")
    ws.append(["项目名称", "单方造价"])
    n = 20000
    for i in range(n):
        ws.append([f"proj{i:06d}", 1000.0 + (i % 50)])
    content = _save_wb(wb)

    df = read_table(content, "big20k.xlsx")
    assert isinstance(df, pd.DataFrame)
    assert len(df) == n, f"期望 {n} 行，实读 {len(df)}（可能被截断）"


def test_import_from_excel_end_to_end_via_fast_path():
    """import_from_excel：快路径读取 → 现有质量清洗 → 入库；异常行被过滤。"""
    wb = Workbook()
    ws = wb.active
    ws.append(REQUIRED)
    ws.append(["测试住宅", "住宅", "框架结构", 12000.0, 18, "四川", 2020, "精装", 5.4e7, 4500.0])
    ws.append(["测试办公", "办公", "剪力墙结构", 8000.0, 12, "北京", 2019, "简装", 4.96e7, 6200.0])
    # 异常行：面积<=0 且 单方造价<500 → 应被质量过滤丢弃
    ws.append(["异常项目", "住宅", "框架结构", 0.0, 5, "上海", 2021, "毛坯", 0.0, 300.0])
    content = _save_wb(wb)

    with tempfile.TemporaryDirectory() as d:
        db_path = os.path.join(d, "cost_prediction.duckdb")
        init_duckdb.init_database(db_path)
        loader = DataLoader(data_dir=d)
        result = loader.import_from_excel(content, "train.xlsx")

        assert result["success"] is True, result
        assert result["imported_count"] == 2, result
        # 异常行被质量校验计入 dropped
        assert result.get("dropped"), result
